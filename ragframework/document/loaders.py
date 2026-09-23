"""Built-in document loaders.

These loaders handle plain text, Markdown, PDF, and DOCX files.
For HTML and other formats see the open issues in
`.github/GOOD_FIRST_ISSUES.md`.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from ragframework.base import Document, DocumentLoader
from ragframework.exceptions import LoaderError

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None


def _make_id(source: str) -> str:
    return hashlib.md5(source.encode()).hexdigest()[:12]


class TextFileLoader(DocumentLoader):
    """Load a plain-text (``.txt``) file as a single :class:`Document`."""

    def __init__(self, encoding: str = "utf-8") -> None:
        self.encoding = encoding

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise LoaderError(f"File not found: {source}")
        if not path.is_file():
            raise LoaderError(f"Not a file: {source}")
        try:
            content = path.read_text(encoding=self.encoding)
        except OSError as exc:
            raise LoaderError(f"Could not read {source}: {exc}") from exc
        return [
            Document(
                id=_make_id(source),
                content=content,
                metadata={"source": source, "filename": path.name},
            )
        ]


class MarkdownLoader(DocumentLoader):
    """Load a Markdown (``.md``) file as a single :class:`Document`.

    The raw Markdown text is stored as-is — no rendering is applied.
    This loader is intentionally minimal; a richer implementation that
    strips front-matter or renders HTML would make a great contribution.
    """

    def __init__(self, encoding: str = "utf-8") -> None:
        self.encoding = encoding

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise LoaderError(f"File not found: {source}")
        if not path.is_file():
            raise LoaderError(f"Not a file: {source}")
        try:
            content = path.read_text(encoding=self.encoding)
        except OSError as exc:
            raise LoaderError(f"Could not read {source}: {exc}") from exc
        return [
            Document(
                id=_make_id(source),
                content=content,
                metadata={"source": source, "filename": path.name, "format": "markdown"},
            )
        ]


class DocxLoader(DocumentLoader):
    """Load a DOCX file into one or more :class:`Document` objects."""

    def __init__(self, split_paragraphs: bool = True) -> None:
        self.split_paragraphs = split_paragraphs

    def load(self, source: str) -> list[Document]:
        path = Path(source)

        if not path.exists():
            raise LoaderError(f"File not found: {source}")

        if not path.is_file():
            raise LoaderError(f"Not a file: {source}")

        if DocxDocument is None:
            raise LoaderError(
                "DOCX support requires 'ragframework[docx]'. "
                "Install it with: pip install ragframework[docx]"
            )

        try:
            docx = DocxDocument(str(path))
        except Exception as exc:
            raise LoaderError(f"Could not read DOCX file {source}: {exc}") from exc
        paragraphs = [paragraph.text for paragraph in docx.paragraphs if paragraph.text.strip()]
        if self.split_paragraphs:
            return [
                Document(
                    id=_make_id(f"{source}_paragraph{i}"),
                    content=paragraph,
                    metadata={
                        "source": source,
                        "filename": path.name,
                        "format": "docx",
                        "paragraph_number": i,
                    },
                )
                for i, paragraph in enumerate(paragraphs, start=1)
            ]

        return [
            Document(
                id=_make_id(source),
                content="\n\n".join(paragraphs),
                metadata={
                    "source": source,
                    "filename": path.name,
                    "format": "docx",
                    "split_paragraphs": False,
                },
            )
        ]


class PDFLoader(DocumentLoader):
    """Load a PDF file into one or more :class:`Document` objects.
    By default, each page is loaded as a separate document with metadata indicating the page number.
    To load the entire PDF as a single document, set ``split_pages=False``
    """

    def __init__(self, split_pages: bool = True) -> None:
        self.split_pages = split_pages

    def load(self, source: str) -> list[Document]:
        path = Path(source)

        if not path.exists():
            raise LoaderError(f"File not found: {source}")
        if not path.is_file():
            raise LoaderError(f"Not a file: {source}")

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ImportError(
                "PDF support requires 'ragframework[pdf]'. "
                "Install it with: pip install ragframework[pdf]"
            ) from exc

        # Open the PDF
        try:
            reader = PdfReader(str(path))
        except Exception as exc:
            raise LoaderError(f"Could not read PDF file {source}: {exc}") from exc

        documents: list[Document] = []

        if not self.split_pages:
            # One single Document for the whole PDF
            try:
                full_text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception as exc:
                raise LoaderError(f"Failed to extract text from {source}: {exc}") from exc

            documents.append(
                Document(
                    id=_make_id(source),
                    content=full_text.strip(),
                    metadata={
                        "source": source,
                        "filename": path.name,
                        "format": "pdf",
                        "total_pages": len(reader.pages),
                        "split_pages": False,
                    },
                )
            )
        else:
            # One Document per page (default)
            for i, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                except Exception as exc:
                    raise LoaderError(
                        f"Failed to extract text from page {i} of {source}: {exc}"
                    ) from exc

                documents.append(
                    Document(
                        id=_make_id(f"{source}_page{i}"),
                        content=text,
                        metadata={
                            "source": source,
                            "filename": path.name,
                            "format": "pdf",
                            "page_number": i,
                            "total_pages": len(reader.pages),
                        },
                    )
                )
        return documents
