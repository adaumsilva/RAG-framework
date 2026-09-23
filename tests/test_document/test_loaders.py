"""Tests for built-in document loaders."""

import builtins
import sys
import types

import pytest

from ragframework.document.loaders import (
    DocxLoader,
    MarkdownLoader,
    PDFLoader,
    TextFileLoader,
)
from ragframework.exceptions import LoaderError


class TestTextFileLoader:
    def test_loads_file(self, tmp_text_file):
        loader = TextFileLoader()
        docs = loader.load(tmp_text_file)
        assert len(docs) == 1
        assert "Hello world" in docs[0].content

    def test_metadata_contains_source(self, tmp_text_file):
        loader = TextFileLoader()
        docs = loader.load(tmp_text_file)
        assert docs[0].metadata["source"] == tmp_text_file

    def test_missing_file_raises(self):
        loader = TextFileLoader()
        with pytest.raises(LoaderError, match="not found"):
            loader.load("/nonexistent/path/file.txt")

    def test_directory_raises(self, tmp_path):
        loader = TextFileLoader()
        with pytest.raises(LoaderError, match="Not a file"):
            loader.load(str(tmp_path))


class TestMarkdownLoader:
    def test_loads_markdown(self, tmp_path):
        md = tmp_path / "doc.md"
        md.write_text("# Title\n\nSome content.", encoding="utf-8")
        loader = MarkdownLoader()
        docs = loader.load(str(md))
        assert len(docs) == 1
        assert "# Title" in docs[0].content
        assert docs[0].metadata["format"] == "markdown"

    def test_missing_file_raises(self):
        loader = MarkdownLoader()
        with pytest.raises(LoaderError):
            loader.load("/no/such/file.md")


class TestDocxLoader:
    def test_loads_paragraphs(self, tmp_path):
        from docx import Document as DocxDocument

        docx_path = tmp_path / "doc.docx"

        doc = DocxDocument()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("")
        doc.add_paragraph("Second paragraph")
        doc.save(docx_path)

        loader = DocxLoader()
        docs = loader.load(str(docx_path))

        assert len(docs) == 2
        assert docs[0].content == "First paragraph"
        assert docs[1].content == "Second paragraph"
        assert docs[0].metadata["format"] == "docx"
        assert docs[0].metadata["paragraph_number"] == 1
        assert docs[1].metadata["paragraph_number"] == 2

    def test_loads_whole_file(self, tmp_path):
        from docx import Document as DocxDocument

        docx_path = tmp_path / "doc.docx"

        doc = DocxDocument()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.save(docx_path)

        loader = DocxLoader(split_paragraphs=False)
        docs = loader.load(str(docx_path))

        assert len(docs) == 1
        assert "First paragraph" in docs[0].content
        assert "Second paragraph" in docs[0].content
        assert docs[0].metadata["format"] == "docx"
        assert docs[0].metadata["split_paragraphs"] is False

    def test_missing_file_raises(self):
        loader = DocxLoader()

        with pytest.raises(LoaderError, match="File not found"):
            loader.load("/nonexistent/path/file.docx")

    def test_invalid_docx_raises(self, tmp_path):
        docx_path = tmp_path / "invalid.docx"
        docx_path.write_text("This is not a valid DOCX file.", encoding="utf-8")

        loader = DocxLoader()

        with pytest.raises(LoaderError, match="Could not read DOCX file"):
            loader.load(str(docx_path))


def test_pdf_loader_requires_pypdf(tmp_path, monkeypatch):
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%EOF\n")

    real_import = builtins.__import__

    def fake_import(name, globals=None, locals=None, fromlist=(), level=0):
        if name == "pypdf":
            raise ImportError("No module named 'pypdf'")
        return real_import(name, globals, locals, fromlist, level)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    loader = PDFLoader()
    with pytest.raises(ImportError, match=r"PDF support requires 'ragframework\[pdf\]'"):
        loader.load(str(pdf_path))


def test_pdf_loader_loads_pages_with_fake_pypdf(tmp_path, monkeypatch):
    class FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class FakePdfReader:
        def __init__(self, source):
            self.pages = [FakePage("page1"), FakePage("page2")]

    fake_pypdf = types.SimpleNamespace(PdfReader=FakePdfReader)
    monkeypatch.setitem(sys.modules, "pypdf", fake_pypdf)

    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%EOF\n")

    loader = PDFLoader(split_pages=True)
    docs = loader.load(str(pdf_path))

    assert len(docs) == 2
    assert docs[0].content == "page1"
    assert docs[1].content == "page2"

    loader_whole = PDFLoader(split_pages=False)
    docs_whole = loader_whole.load(str(pdf_path))
    assert len(docs_whole) == 1
    assert "page1" in docs_whole[0].content
    assert "page2" in docs_whole[0].content
